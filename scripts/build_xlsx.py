"""data/*.csv -> stock.xlsx (freezer, 4 sheets) and mouse.xlsx (KPC mouse). Run by GitHub Actions on every data change."""
import csv
import datetime as dt
import pathlib

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
BOOKS = {
    "stock.xlsx": [("celline", "celline.csv"), ("bacteria", "bacteria.csv"),
                   ("reagent", "reagent.csv"), ("siRNA", "siRNA.csv")],
    "mouse.xlsx": [("KPC mouse", "mice.csv")],
}
DATE_COLS = {"date", "dob", "weaning", "genotyping_date"}
NUMERIC_COLS = {"concentration (ng/µl)", "generation"}  # pure numbers; other concentrations keep their unit text
HEADER_FILL = PatternFill("solid", fgColor="DAE6F0")
# same colors as the original KPC Excel conditional format (M blue, F pink)
SEX_FILL = {"M": PatternFill(bgColor="D9E1F2"), "F": PatternFill(bgColor="F5DDE4")}


def convert(col, value):
    if value == "":
        return None
    if col in DATE_COLS:
        try:
            return dt.datetime.strptime(value, "%Y-%m-%d").date()
        except ValueError:
            return value
    if col in NUMERIC_COLS:
        try:
            return float(value) if "." in value else int(value)
        except ValueError:
            return value
    return value


def add_sheet(wb, name, fname):
    ws = wb.create_sheet(name)
    with open(ROOT / "data" / fname, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    header = rows[0]
    ws.append(header)
    for r in rows[1:]:
        r = r + [""] * (len(header) - len(r))
        ws.append([convert(c, v) for c, v in zip(header, r)])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for ci, col in enumerate(header, start=1):
        letter = get_column_letter(ci)
        for (cell,) in ws.iter_rows(min_row=2, min_col=ci, max_col=ci):
            if col in DATE_COLS and isinstance(cell.value, dt.date):
                cell.number_format = "yyyy-mm-dd"
            elif col not in NUMERIC_COLS:
                cell.number_format = "@"
        width = max([len(str(col))] + [len(str(r[ci - 1])) for r in rows[1:] if len(r) >= ci])
        ws.column_dimensions[letter].width = min(max(width + 2, 8), 60)
    if "sex" in header:
        letter = get_column_letter(header.index("sex") + 1)
        rng = f"{letter}2:{letter}{max(len(rows), 2) + 500}"  # leave room for rows added later in Excel
        for val, fill in SEX_FILL.items():
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'UPPER({letter}2)="{val}"'], fill=fill))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    for book, sheets in BOOKS.items():
        wb = Workbook()
        wb.remove(wb.active)
        for name, fname in sheets:
            if (ROOT / "data" / fname).exists():
                add_sheet(wb, name, fname)
        wb.save(ROOT / book)
        print(f"{book} written")


if __name__ == "__main__":
    main()
